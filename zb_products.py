# -*- coding: utf-8 -*-
"""
徕卡产品销售分类

从 ERP 合同的产品明细里，把卖的相机/镜头/配件分门别类：
  · 品类：相机 / 镜头 / 配件 / 周边 / 激光电视
  · 相机系列：M 系列 / Q 系列 / SL 系列 / D-LUX / 拍立得 …
  · 具体型号：M11-P、Q3、D-LUX 8、SOFORT 2 …
  · 品牌归属：徕卡原厂 / 徕卡副厂 / 其他品牌

分类规则全部在 config.json 的「产品分类」里，改配置即可，不用改代码。
"""
from __future__ import annotations

import re
from collections import defaultdict

import zbom_core as core

# 型号识别：越具体的写前面，避免 M11 抢了 M11-P。
# 注意：不能用 \b 做边界——中文是 \w，'M9银' 会被判成没有边界，导致漏匹配。
NB = r"(?![0-9A-Za-z])"          # 后面不能再跟字母数字
MODEL_RULES: list[tuple[str, str]] = [
    # 拍立得
    (r"SOFORT\s*2", "SOFORT 2"),
    (r"SOFORT", "SOFORT"),
    (r"拍立得", "拍立得"),
    # Q 系列
    (r"Q3\s*43", "Q3 43"),
    (r"Q3\s*28.*黑白|Q3\s*28\s*[-–]?\s*M", "Q3 28 Monochrom"),
    (r"Q3\s*28", "Q3 28"),
    (r"Q3" + NB, "Q3"),
    (r"Q2.*黑白|Q2\s*Monochrom", "Q2 Monochrom"),
    (r"Q2" + NB, "Q2"),
    (r"Q\s*\(Typ\s*116\)", "Q (Typ 116)"),
    # M 系列
    (r"M11\s*[-–]?\s*P", "M11-P"),
    (r"M11\s*Monochrom", "M11 Monochrom"),
    (r"M11" + NB, "M11"),
    (r"M10\s*[-–]?\s*R", "M10-R"),
    (r"M10\s*[-–]?\s*P", "M10-P"),
    (r"M10\s*Monochrom", "M10 Monochrom"),
    (r"M10" + NB, "M10"),
    (r"M9" + NB, "M9"),
    (r"M8" + NB, "M8"),
    (r"M7" + NB, "M7"),
    (r"M6" + NB, "M6"),
    (r"M240" + NB, "M240"),
    (r"M\s*EV1", "M EV1"),
    (r"MP\s*0\.72|MP" + NB, "MP"),
    (r"M-P" + NB, "M-P"),
    (r"M-A" + NB, "M-A"),
    (r"X-U" + NB, "X-U"),
    (r"X1" + NB, "X1"),
    (r"X2" + NB, "X2"),
    (r"相机\s*T\s+黑|T\s*\(Typ", "T (Typ 701)"),
    # SL / S
    (r"SL3\s*[-–]?\s*P", "SL3-P"),
    (r"SL3" + NB, "SL3"),
    (r"SL2\s*[-–]?\s*S", "SL2-S"),
    (r"SL2" + NB, "SL2"),
    (r"S3" + NB, "S3"),
    (r"S2" + NB, "S2"),
    # 便携 / 其他
    (r"D\s*[-–]?\s*LUX\s*8", "D-LUX 8"),
    (r"D\s*[-–]?\s*LUX\s*7", "D-LUX 7"),
    (r"D\s*[-–]?\s*LUX", "D-LUX"),
    (r"V\s*[-–]?\s*LUX\s*5", "V-LUX 5"),
    (r"V\s*[-–]?\s*LUX", "V-LUX"),
    (r"C\s*[-–]?\s*LUX", "C-LUX"),
    (r"Cine\s*Play", "Cine Play 1"),
    (r"Cine\s*1|激光电视", "Cine 1"),
]

# 镜头焦段：50mm/1.4 这种
LENS_RE = re.compile(r"(\d{2,3})\s*mm\s*/?\s*(?:f/?)?\s*(0\.95|1\.25|1\.4|1\.2|2\.8|2|3\.5|4|5\.6)", re.I)


def _match_rule(title: str, rules: list[dict]) -> str:
    for r in rules or []:
        kws = r.get("关键词") or []
        if not kws:
            continue
        for kw in kws:
            if kw and kw in title:
                return r.get("名称") or ""
    # 没有关键词的兜底规则（关键词为空数组）
    for r in rules or []:
        if not (r.get("关键词") or []):
            return r.get("名称") or ""
    return ""


def extract_model(title: str) -> str:
    t = title or ""
    for pat, name in MODEL_RULES:
        if re.search(pat, t, re.I):
            return name
    m = LENS_RE.search(t)
    if m:
        return f"{m.group(1)}mm f/{m.group(2)}"
    return ""


def extract_goods_no(title: str) -> str:
    """产品名结尾一般是货号，如 '... 20071' 或 '... CP_20251204009'"""
    m = re.search(r"\s([A-Za-z]{0,3}_?[0-9A-Za-z]{4,16})\s*$", title or "")
    return m.group(1) if m else ""


def clean_name(title: str) -> str:
    """非相机/镜头类产品：去掉前缀和结尾货号，截断当型号用"""
    s = re.sub(r"^(自用|李和林已售|陈列)\s*", "", core.cell_text(title))
    s = re.sub(r"^(二手|全新|样机)\s*", "", s)
    s = re.sub(r"^(徕卡|Leica)\s*", "", s)
    s = re.sub(r"\s*[A-Za-z]{0,3}_?[0-9A-Za-z]{4,16}\s*/?\s*$", "", s)
    s = re.sub(r"[/／]\s*(准新|\d{1,2}新|9新)\s*$", "", s)   # 去掉「/ 99新」这类尾注
    s = re.sub(r"[/／\s]+$", "", s)
    return s.strip()[:26] or "未识别"


# --------------------------------------------------------------------------
# 型号拆解口径：相机按系列、镜头按卡口系列、其余按品类
#   ERP 里的产品名形如「徕卡镜头 M 35mm/2 AA」「徕卡镜头 SL 24-70/2.8」，
#   「镜头」后面那个字母就是卡口系列（M / SL / TL），据此把镜头分到对应系列。
# --------------------------------------------------------------------------

_LENS_SERIES_RE = re.compile(r"镜头\s*([A-Za-z]{1,3})\b")


def lens_series(title: str) -> str:
    """从产品名里取镜头的卡口系列，取不到返回空串"""
    m = _LENS_SERIES_RE.search(core.cell_text(title) or "")
    if not m:
        return ""
    s = m.group(1).upper()
    # 别把「徕卡镜头 35mm」里的数字或杂词当系列
    return s if s in ("M", "SL", "TL", "S", "CL") else ""


def breakdown_cat(item: dict, cfg: dict) -> str:
    """把一条产品记录归到「型号拆解」用的大类。

    规则写在 config.json 的「型号拆解分类」里，想调口径改配置即可，不用改代码。
    """
    rules = (cfg.get("产品分类", {}) or {}).get("型号拆解分类", {}) or {}
    fallback = rules.get("兜底") or "其他"
    cat = item.get("品类") or ""

    if cat == "相机":
        # 相机单独兜底，别落到全局「其他」——「其他」里还混着投影电视等非相机，
        # 混在一起会让「相机占比」算高。
        return ((rules.get("相机系列映射") or {}).get(item.get("系列") or "")
                or rules.get("相机兜底") or fallback)
    if cat == "镜头":
        key = lens_series(item.get("产品名") or "")
        if key:
            mapped = (rules.get("镜头系列映射") or {}).get(key)
            if mapped:
                return mapped
        return fallback
    return (rules.get("品类映射") or {}).get(cat, fallback)


def classify(title: str, cfg: dict) -> dict:
    """给一个产品名打上 品类/系列/型号/品牌/新旧"""
    t = core.cell_text(title)
    rules = cfg.get("产品分类", {}) or {}

    category = _match_rule(t, rules.get("品类规则")) or "其他"
    brand = _match_rule(t, rules.get("品牌规则")) or ("徕卡原厂" if "徕卡" in t else "其他品牌")
    condition = _match_rule(t, rules.get("新旧规则")) or "全新"

    series = ""
    if category == "相机":
        series = _match_rule(t, rules.get("相机系列")) or "其他相机"
    elif category == "镜头":
        series = "镜头"

    # 型号只在相机/镜头里有意义；配件周边用清理过的产品名，免得把 D-LUX 之类带进来
    if category == "相机":
        model = extract_model(t) or clean_name(t)
    elif category == "镜头":
        m = LENS_RE.search(t)
        model = f"{m.group(1)}mm f/{m.group(2)}" if m else clean_name(t)
    else:
        model = clean_name(t)

    return {"品类": category, "品牌": brand, "系列": series, "型号": model,
            "新旧": condition, "货号": extract_goods_no(t), "产品名": t}


def classify_all(cache: dict, contracts: list[dict], cfg: dict) -> list[dict]:
    """把 {合同ord: [产品行]} 展开成一维的产品销售记录"""
    meta = {str(c["ord"]): c for c in contracts}
    out: list[dict] = []
    for ord_, prods in (cache or {}).items():
        c = meta.get(str(ord_))
        if not c:
            continue
        for p in prods or []:
            title = core.cell_text(p.get("title"))
            if not title:
                continue
            qty = core.to_money(p.get("num1")) or 1.0
            amount = core.to_money(p.get("money1"))
            info = classify(title, cfg)
            out.append({
                **info,
                "ord": str(ord_),
                "htid": core.cell_text(c.get("htid")),
                "门店": core.cell_text(c.get("htfl")),
                "销售员": _person(c.get("xsry")),
                "日期": (core.to_date(c.get("dateQD")) or "").isoformat() if core.to_date(c.get("dateQD")) else "",
                "合同金额": core.to_money(c.get("moneyall")),
                "数量": qty,
                "金额": amount,
            })
    return out


def _person(raw) -> str:
    s = core.cell_text(raw)
    for sep in ("-", "－", "—"):
        if sep in s:
            head, _, tail = s.partition(sep)
            if tail.strip():
                return tail.strip()
    return s


# --------------------------------------------------------------------------
# 汇总
# --------------------------------------------------------------------------

def summarize(items: list[dict], scope_stores: set[str] | None = None) -> dict:
    if scope_stores:
        items = [x for x in items if x["门店"] in scope_stores]

    total_amount = sum(x["金额"] for x in items)
    total_qty = sum(x["数量"] for x in items)
    cams = [x for x in items if x["品类"] == "相机"]
    lenses = [x for x in items if x["品类"] == "镜头"]

    return {
        "总计": {"金额": round(total_amount, 2), "数量": round(total_qty, 2),
                 "产品行": len(items), "合同数": len({x["ord"] for x in items})},
        "按品类": summarize_group(items, "品类"),
        "按品牌": summarize_group(items, "品牌"),
        "按新旧": summarize_group(items, "新旧"),
        "相机系列": summarize_group(cams, "系列"),
        "相机型号": summarize_group(cams, "型号"),
        "相机新旧": summarize_group(cams, "新旧"),
        "镜头型号": summarize_group(lenses, "型号"),
        "全部型号": summarize_group(items, "型号"),
        "按门店": summarize_group(items, "门店"),
        "按销售员": summarize_group(items, "销售员"),
        "明细": items,
    }


def summarize_group(items: list[dict], key: str) -> list[dict]:
    g: dict[str, dict] = defaultdict(lambda: {"数量": 0.0, "金额": 0.0, "合同": set()})
    for it in items:
        k = it.get(key) or "（空）"
        g[k]["数量"] += it["数量"]
        g[k]["金额"] += it["金额"]
        g[k]["合同"].add(it["ord"])
    total = sum(v["金额"] for v in g.values()) or 1
    rows = [{"名称": k, "数量": round(v["数量"], 2), "金额": round(v["金额"], 2),
             "合同数": len(v["合同"]), "占比": v["金额"] / total} for k, v in g.items()]
    rows.sort(key=lambda r: -r["金额"])
    return rows


# --------------------------------------------------------------------------
# Excel 报表
# --------------------------------------------------------------------------

def write_report(cfg: dict, items: list[dict], scope: str = "商场店", tag: str = "") -> "Path":
    import datetime as _dt
    from pathlib import Path

    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    THIN = Side(style="thin", color="BFBFBF")
    BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
    HEAD_FILL = PatternFill("solid", fgColor="1F4E78")
    HEAD_FONT = Font(name="微软雅黑", size=10, bold=True, color="FFFFFF")
    TITLE_FONT = Font(name="微软雅黑", size=14, bold=True, color="1F4E78")
    SUB_FONT = Font(name="微软雅黑", size=9, color="808080")
    BODY_FONT = Font(name="微软雅黑", size=10)
    BOLD = Font(name="微软雅黑", size=10, bold=True)
    GOLD = PatternFill("solid", fgColor="FFF2CC")
    SILVER = PatternFill("solid", fgColor="F2F2F2")
    BRONZE = PatternFill("solid", fgColor="FCE4D6")
    MONEY = '#,##0.00'
    PCT = '0.0%'
    CENTER = Alignment(horizontal="center", vertical="center")
    LEFT = Alignment(horizontal="left", vertical="center")

    stores_cfg = cfg.get("门店", {}) or {}
    if scope == "商场店":
        scope_stores = {n for n, s in stores_cfg.items()
                        if isinstance(s, dict) and (s.get("类型") or "") == "商场店"}
    else:
        scope_stores = None

    s = summarize(items, scope_stores)
    today = _dt.datetime.now().strftime("%Y-%m-%d %H:%M")
    wb = Workbook()

    def head(ws, row, headers, widths):
        for i, h in enumerate(headers, 1):
            c = ws.cell(row=row, column=i, value=h)
            c.fill, c.font, c.alignment, c.border = HEAD_FILL, HEAD_FONT, CENTER, BORDER
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w

    def title(ws, text, sub, span):
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=span)
        ws.cell(row=1, column=1, value=text).font = TITLE_FONT
        ws.row_dimensions[1].height = 26
        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=span)
        ws.cell(row=2, column=1, value=sub).font = SUB_FONT

    def table(ws, start, rows, name_head="名称"):
        head(ws, start, [name_head, "数量", "金额", "涉及合同数", "金额占比"],
             [30, 10, 16, 12, 11])
        r = start + 1
        for i, row in enumerate(rows, 1):
            for j, v in enumerate([row["名称"], row["数量"], row["金额"], row["合同数"], row["占比"]], 1):
                c = ws.cell(row=r, column=j, value=v)
                c.font, c.border = BODY_FONT, BORDER
                c.alignment = LEFT if j == 1 else CENTER
            ws.cell(row=r, column=3).number_format = MONEY
            ws.cell(row=r, column=5).number_format = PCT
            f = {1: GOLD, 2: SILVER, 3: BRONZE}.get(i)
            if f:
                for j in range(1, 6):
                    ws.cell(row=r, column=j).fill = f
            r += 1
        tot = ws.cell(row=r, column=1, value="合计")
        tot.font = BOLD
        ws.cell(row=r, column=2, value=round(sum(x["数量"] for x in rows), 2)).font = BOLD
        c = ws.cell(row=r, column=3, value=round(sum(x["金额"] for x in rows), 2))
        c.font, c.number_format = BOLD, MONEY
        ws.cell(row=r, column=5, value=1).number_format = PCT
        return r + 2

    # ---------------- 相机（重点）----------------
    ws = wb.active
    ws.title = "徕卡相机分类"
    title(ws, "徕卡相机销售分类", f"范围：{scope}　|　口径：合同金额（优惠后）　|　生成：{today}", 5)
    r = 4
    ws.cell(row=r, column=1, value="按系列").font = TITLE_FONT
    r = table(ws, r + 1, s["相机系列"], "相机系列")
    ws.cell(row=r, column=1, value="按型号").font = TITLE_FONT
    r = table(ws, r + 1, s["相机型号"], "相机型号")

    # ---------------- 全品类 ----------------
    ws2 = wb.create_sheet("品类与品牌")
    title(ws2, "产品品类 / 品牌结构", f"范围：{scope}　|　生成：{today}", 5)
    r = 4
    ws2.cell(row=r, column=1, value="按品类").font = TITLE_FONT
    r = table(ws2, r + 1, s["按品类"], "品类")
    ws2.cell(row=r, column=1, value="按品牌归属").font = TITLE_FONT
    r = table(ws2, r + 1, s["按品牌"], "品牌")

    # ---------------- 镜头 & 配件 ----------------
    ws3 = wb.create_sheet("镜头与配件")
    title(ws3, "镜头 / 配件 / 周边 明细分类", f"范围：{scope}　|　生成：{today}", 5)
    r = 4
    ws3.cell(row=r, column=1, value="镜头型号").font = TITLE_FONT
    r = table(ws3, r + 1, s["镜头型号"], "镜头型号")
    # 配件/周边/电视：从全部型号里挑出属于这些品类的（用明细反查品类）
    model_cat: dict[str, set] = defaultdict(set)
    for it in s["明细"]:
        model_cat[it["型号"]].add(it["品类"])
    pick = [x for x in s["全部型号"] if model_cat.get(x["名称"], set()) & {"配件", "周边", "激光电视"}]
    ws3.cell(row=r, column=1, value="配件 / 周边 / 电视").font = TITLE_FONT
    table(ws3, r + 1, pick[:60], "产品")

    # ---------------- 门店对比 ----------------
    ws4 = wb.create_sheet("门店对比")
    title(ws4, "各门店卖的相机", f"范围：{scope}　|　生成：{today}", 6)
    head(ws4, 4, ["门店", "总金额", "相机金额", "相机台数", "相机占比", "热销型号"],
         [16, 16, 16, 12, 11, 34])
    r = 5
    by_store = {x["名称"]: x for x in s["按门店"]}
    for store in (scope_stores if scope_stores else by_store.keys()):
        store_items = [x for x in s["明细"] if x["门店"] == store]
        cams = [x for x in store_items if x["品类"] == "相机"]
        amt = sum(x["金额"] for x in store_items)
        cam_amt = sum(x["金额"] for x in cams)
        top = summarize_group(cams, "型号")[:3]
        vals = [store, amt, cam_amt, sum(x["数量"] for x in cams),
                (cam_amt / amt) if amt else 0,
                "、".join(f"{t['名称']}({t['数量']:.0f})" for t in top)]
        for j, v in enumerate(vals, 1):
            c = ws4.cell(row=r, column=j, value=v)
            c.font, c.border = BODY_FONT, BORDER
            c.alignment = LEFT if j in (1, 6) else CENTER
        for j in (2, 3):
            ws4.cell(row=r, column=j).number_format = MONEY
        ws4.cell(row=r, column=5).number_format = PCT
        r += 1

    # ---------------- 明细 ----------------
    ws5 = wb.create_sheet("产品明细")
    title(ws5, "产品销售明细", f"共 {len(s['明细'])} 行　|　生成：{today}", 11)
    head(ws5, 4, ["签订日期", "门店", "销售员", "合同编号", "品类", "相机系列", "型号",
                  "品牌", "数量", "金额", "产品全名"],
         [12, 13, 10, 18, 9, 11, 14, 11, 7, 12, 46])
    r = 5
    for it in sorted(s["明细"], key=lambda x: (x["日期"], x["门店"]), reverse=True):
        for j, v in enumerate([it["日期"], it["门店"], it["销售员"], it["htid"], it["品类"],
                               it["系列"], it["型号"], it["品牌"], it["数量"], it["金额"],
                               it["产品名"]], 1):
            c = ws5.cell(row=r, column=j, value=v)
            c.font, c.border = BODY_FONT, BORDER
            c.alignment = LEFT if j in (4, 11, 3, 2) else CENTER
        ws5.cell(row=r, column=10).number_format = MONEY
        r += 1

    # ---------------- 保存 ----------------
    from pathlib import Path as _P
    out_dir = core.cfg_path(cfg, "目录")
    out_dir.mkdir(parents=True, exist_ok=True)
    stamp = _dt.datetime.now().strftime("%Y%m%d_%H%M")
    target = out_dir / f"产品分类_{scope}_{tag + '_' if tag else ''}{stamp}.xlsx"
    wb.save(target)
    return target


# --------------------------------------------------------------------------
# 命令行
# --------------------------------------------------------------------------

def load_items(cfg: dict, tag: str = "", log=print):
    """读缓存 + 合同列表，产出分类后的产品记录"""
    import json

    from zb_erp import PRODUCT_CACHE_REL, ZbErp

    cache_path = core.BASE / PRODUCT_CACHE_REL
    if not cache_path.exists():
        raise SystemExit(
            f"还没有产品明细缓存：{cache_path}\n"
            "先跑： python zb_products.py --抓取"
        )
    cache = json.loads(cache_path.read_text(encoding="utf-8"))

    e = cfg.get("智邦ERP") or {}
    dr = e.get("日期范围") or {}
    d_from, d_to = dr.get("开始") or "", dr.get("结束") or ""
    erp = ZbErp(e["地址"], e["账号"], e["密码"], e.get("串号") or "zb-scraper-01")
    erp.login(log=lambda *_: None)
    contracts = erp.contracts(d_from, d_to, page_size=300, log=None)
    log(f"  合同 {len(contracts)} 份，产品缓存 {len(cache)} 份")
    items = classify_all(cache, contracts, cfg)
    log(f"  产品行 {len(items)} 条")
    return items, contracts, d_from, d_to


def main(argv=None) -> int:
    import argparse
    import sys

    ap = argparse.ArgumentParser(description="徕卡产品销售分类")
    ap.add_argument("--抓取", action="store_true", help="先抓产品明细（单线程限速，几分钟）")
    ap.add_argument("--出报表", action="store_true", help="生成 Excel 报表")
    ap.add_argument("--范围", default="商场店", help="商场店 / 全部")
    args = ap.parse_args(argv)

    cfg = core.load_config()
    e = cfg.get("智邦ERP") or {}

    if args.抓取:
        from zb_erp import ErpSessionDown, ZbErp
        dr = e.get("日期范围") or {}
        d_from, d_to = dr.get("开始") or "", dr.get("结束") or ""
        erp = ZbErp(e["地址"], e["账号"], e["密码"], e.get("串号") or "zb-scraper-01")
        erp.login()
        contracts = erp.contracts(d_from, d_to, page_size=300, log=None)
        print(f"合同 {len(contracts)} 份，开始温柔抓取产品明细（单线程 + 限速）…")
        try:
            erp.all_products(contracts, delay=0.35)
        except ErpSessionDown as exc:
            print("\n[中止] " + str(exc), file=sys.stderr)
            return 2
        if not args.出报表:
            return 0

    items, contracts, d_from, d_to = load_items(cfg)

    stores_cfg = cfg.get("门店", {}) or {}
    scope_stores = None
    if args.范围 == "商场店":
        scope_stores = {n for n, s in stores_cfg.items()
                        if isinstance(s, dict) and (s.get("类型") or "") == "商场店"}
    s = summarize(items, scope_stores)

    print(f"\n{'='*64}\n产品分类汇总（范围：{args.范围}）\n{'='*64}")
    t = s["总计"]
    print(f"产品行 {t['产品行']} 条 / 涉及 {t['合同数']} 份合同 / 金额 {t['金额']:,.2f} 元")

    print("\n【按品类】")
    for r in s["按品类"]:
        print(f"  {r['名称']:<8} {r['数量']:>7.0f} 件  {r['金额']:>15,.2f} 元  {r['占比']*100:5.1f}%")

    print("\n【相机按系列】")
    for r in s["相机系列"]:
        print(f"  {r['名称']:<12} {r['数量']:>5.0f} 台  {r['金额']:>15,.2f} 元  {r['占比']*100:5.1f}%")

    print("\n【相机 全新 / 二手】")
    for r in s["相机新旧"]:
        print(f"  {r['名称']:<6} {r['数量']:>5.0f} 台  {r['金额']:>15,.2f} 元  {r['占比']*100:5.1f}%")

    print("\n【相机型号排行】")
    for i, r in enumerate(s["相机型号"], 1):
        print(f"  {i:>2}. {r['名称']:<18} {r['数量']:>5.0f} 台  {r['金额']:>15,.2f} 元  {r['占比']*100:5.1f}%")

    if args.出报表:
        target = write_report(cfg, items, scope=args.范围,
                              tag=f"{d_from}_{d_to}" if d_from or d_to else "")
        print(f"\n报表已生成：\n  {target}")
    return 0


if __name__ == "__main__":
    import sys
    sys.exit(main())

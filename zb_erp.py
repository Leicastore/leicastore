# -*- coding: utf-8 -*-
"""
智邦国际 ERP（ZBIntel）业绩抓取客户端

你们公司用的 ERP 是「智邦国际」，不是志邦家居。系统地址由动态域名提供。

接口要点（都是从系统自带的 OpenAPI 文档里逆出来的）：
  1. 所有接口都是 POST + JSON
  2. 登录 /webapi/v3/ov1/login，参数包在 datas 数组里：{"datas":[{"id":"user","val":"txt:xxx"}, ...]}
     - 明文必须加前缀 "txt:"，否则服务端判定为空
     - user / password / serialnum 三个都必填，serialnum 随便一串 20 位以内字母数字
     - 登录失败超过 3 次会要求 rndcode 验证码
  3. 登录成功返回 header.session，之后：
     - v3 原生接口：session 放 HTTP 头 ZBAPI-Token
     - ASP 老接口：session 放 body 里，且要带 cmdkey（每个接口固定值，合同列表是 "refresh"）
  4. ASP 接口必须直接调 .asp 地址，走 /webapi/v3/ov1/ 别名会返回空
  5. 合同列表返回 body.source.table.{page, cols, rows}

用法：
    python zb_erp.py                      # 按 config.json 抓取
    python zb_erp.py --开始 2026-01-01 --结束 2026-12-31
    python zb_erp.py --出报表              # 抓完直接生成 Excel 报表
    python zb_erp.py --探测                # 只看接口和数据结构，不写文件
"""
from __future__ import annotations

import argparse
import json
import ssl
import sys
import time
import urllib.error
import urllib.parse
import urllib.request
from pathlib import Path

import zbom_core as core

UA = "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/124.0 Safari/537.36"

# 合同列表接口的固定 cmdkey（来自接口文档的 aspfixcmdkey 字段）
CMDKEY_CONTRACT = "refresh"
PATH_CONTRACT = "/sysa/mobilephone/salesmanage/contract/billlist.asp"


class ErpError(RuntimeError):
    pass


class ErpSessionDown(ErpError):
    """ERP 自己的会话存储（Redis）不可用。这是服务端故障，重试没用，要等它恢复。"""

    def __init__(self, msg: str):
        super().__init__(
            "ERP 服务端会话异常（它自己的 Redis 会话存储挂了）：\n"
            f"  {msg}\n"
            "→ 这是对方服务器的问题，不是本工具的问题。\n"
            "→ 已经停止请求。等几分钟再试；如果一直这样，联系 ERP 管理员重启会话服务。"
        )


# 合同详情接口（带产品明细）。这个接口不需要 cmdkey。
PATH_BILL_DETAIL = "/sysa/mobilephone/salesmanage/contract/BillService.asp"
# 产品明细缓存（断点续抓，避免重复请求把 ERP 压垮）
PRODUCT_CACHE_REL = "数据/缓存/合同产品明细.json"


class ZbErp:
    def __init__(self, base: str, account: str, password: str, serial: str = "zb-scraper-01",
                 timeout: int = 90):
        self.base = base.rstrip("/")
        self.account = account
        self.password = password
        self.serial = serial
        self.timeout = timeout
        self.session: str | None = None
        self.cookies: dict[str, str] = {}
        self._ctx = ssl.create_default_context()
        self._ctx.check_hostname = False
        self._ctx.verify_mode = ssl.CERT_NONE

    # ---------------------------------------------------------------- 底层
    def _post(self, path: str, body: dict, token: str | None = None) -> str:
        url = self.base + path
        data = json.dumps(body, ensure_ascii=False).encode("utf-8")
        req = urllib.request.Request(url, data=data, method="POST")
        req.add_header("Content-Type", "application/json;charset=UTF-8")
        req.add_header("User-Agent", UA)
        if self.cookies:
            req.add_header("Cookie", "; ".join(f"{k}={v}" for k, v in self.cookies.items()))
        if token:
            req.add_header("ZBAPI-Token", token)
        try:
            with urllib.request.urlopen(req, timeout=self.timeout, context=self._ctx) as resp:
                for sc in resp.headers.get_all("Set-Cookie") or []:
                    kv = sc.split(";")[0]
                    if "=" in kv:
                        k, v = kv.split("=", 1)
                        if v:
                            self.cookies[k.strip()] = v.strip()
                return resp.read().decode("utf-8", "replace")
        except urllib.error.HTTPError as e:
            raise ErpError(f"HTTP {e.code} {e.reason}  ({url})") from None
        except urllib.error.URLError as e:
            raise ErpError(
                f"连不上 {url}\n原因：{e.reason}\n"
                "· 这个 ERP 走的是动态域名，需要能上外网\n"
                "· 用浏览器先打开一次确认地址有效"
            ) from None

    # ---------------------------------------------------------------- 登录
    def login(self, log=print) -> dict:
        body = {"datas": [
            {"id": "user", "val": f"txt:{self.account}"},
            {"id": "password", "val": f"txt:{self.password}"},
            {"id": "serialnum", "val": f"txt:{self.serial}"},
        ]}
        out = self._post("/webapi/v3/ov1/login", body)
        try:
            d = json.loads(out)
        except json.JSONDecodeError:
            raise ErpError(f"登录返回的不是 JSON：{out[:300]}") from None

        status = d.get("header", {}).get("status")
        msg = d.get("header", {}).get("message", "")
        if status != 0:
            hint = ""
            if "验证码" in msg or "rndcode" in msg:
                hint = ("\n→ 密码连续输错 3 次，系统要求验证码了。\n"
                        "  用浏览器登录一次系统，把验证码填对，之后再跑脚本。")
            elif "密码错误" in msg:
                hint = ("\n→ 账号或密码不对。注意密码不要带公司前缀，账号是登录名不是员工姓名。")
            raise ErpError(f"登录失败：{msg}{hint}")

        self.session = d["header"].get("session")
        if not self.session:
            raise ErpError("登录成功但没返回 session，无法继续。")
        log(f"  登录成功：{self.account}（session {self.session[:12]}…）")
        return d

    # ---------------------------------------------------------------- 业务接口
    def asp_call(self, path: str, cmdkey: str, params: dict, log=None) -> dict:
        """调用 ASP 老接口。params 是 {参数名: 值}，会自动包成 datas 数组。"""
        if not self.session:
            raise ErpError("还没登录，先调用 login()")
        body = {
            "session": self.session,
            "datas": [{"id": k, "val": "" if v is None else str(v)} for k, v in params.items()],
        }
        if cmdkey:                       # 有些接口没有 cmdkey，传空就别带这个字段
            body["cmdkey"] = cmdkey
        out = self._post(path, body)
        if not out.strip():
            raise ErpError(
                f"接口返回空内容：{path}\n"
                "→ 多半是 session 过期或接口地址不对。重新登录试试。"
            )
        try:
            # strict=False：ERP 返回的 HTML 片段里可能夹着裸控制字符
            d = json.loads(out, strict=False)
        except json.JSONDecodeError:
            raise ErpError(f"接口返回的不是 JSON：{out[:300]}") from None
        st = d.get("header", {}).get("status")
        msg = str(d.get("header", {}).get("message") or "")
        if st == 10 or "会话异常" in msg:
            # ERP 自己的会话存储（Redis）不可用，不是我们的问题，别硬刚
            raise ErpSessionDown(msg or "会话异常")
        if st not in (0, 200, None):
            raise ErpError(f"接口报错：{msg}")
        return d

    def v3_call(self, path: str, params: dict) -> dict:
        """调用 v3 原生接口，session 走 ZBAPI-Token 头"""
        if not self.session:
            raise ErpError("还没登录，先调用 login()")
        out = self._post(path, params, token=self.session)
        try:
            d = json.loads(out)
        except json.JSONDecodeError:
            raise ErpError(f"接口返回的不是 JSON：{out[:300]}") from None
        if d.get("Code") not in (200, None):
            if d.get("Code") == 400:
                raise ErpError("token 验证失败（Code 400），session 过期了，重新登录。")
            raise ErpError(f"接口报错 Code={d.get('Code')} Msg={d.get('Msg')}")
        return d

    def departments(self) -> list[dict]:
        d = self.v3_call("/webapi/v3/orgs/dept/list", {"page_size": 500, "page_index": 1})
        cols = d.get("Cols") or []
        return [dict(zip(cols, r)) for r in (d.get("Rows") or [])]

    # ---------------------------------------------------------- 合同产品明细
    def bill_products(self, ord_) -> list[dict]:
        """取一份合同的产品明细行"""
        d = self.asp_call(PATH_BILL_DETAIL, "", {"ord": ord_})
        groups = (d.get("body", {}).get("bill") or {}).get("groups") or []
        for g in groups:
            if g.get("id") == "productList":
                fields = g.get("fields") or []
                if not fields:
                    return []
                tb = (fields[0].get("source") or {}).get("table") or {}
                cols = [c["id"] if isinstance(c, dict) else c for c in (tb.get("cols") or [])]
                return [dict(zip(cols, r)) for r in (tb.get("rows") or [])]
        return []

    def all_products(self, contracts: list[dict], delay: float = 0.35, workers: int = 1,
                     cache_path: Path | None = None, log=print, max_retry: int = 3) -> dict:
        """
        批量取所有合同的产品明细。

        有意做得**很温柔**：默认单线程 + 每份之间 sleep，因为实测并发猛拉会把
        ERP 的会话服务（Redis）打挂。带断点缓存，中断了下次接着跑。
        """
        cache_path = Path(cache_path) if cache_path else (core.BASE / PRODUCT_CACHE_REL)
        cache_path.parent.mkdir(parents=True, exist_ok=True)
        cache: dict[str, list] = {}
        if cache_path.exists():
            try:
                cache = json.loads(cache_path.read_text(encoding="utf-8"))
                log(f"  断点缓存：已有 {len(cache)} 份合同的产品明细")
            except Exception:
                cache = {}

        todo = [c for c in contracts if str(c["ord"]) not in cache]
        log(f"  需要抓取 {len(todo)} 份（共 {len(contracts)} 份）")

        done = 0
        for c in todo:
            ord_ = str(c["ord"])
            for attempt in range(max_retry):
                try:
                    cache[ord_] = self.bill_products(ord_)
                    break
                except ErpSessionDown:
                    # 服务端会话挂了，保存进度并立刻停手
                    cache_path.write_text(json.dumps(cache, ensure_ascii=False), encoding="utf-8")
                    raise
                except Exception as exc:
                    if attempt == max_retry - 1:
                        log(f"    ord={ord_} 放弃：{exc}")
                        cache[ord_] = []
                    else:
                        time.sleep(1.5 * (attempt + 1))
            done += 1
            if done % 50 == 0:
                cache_path.write_text(json.dumps(cache, ensure_ascii=False), encoding="utf-8")
                log(f"    进度 {done}/{len(todo)}（已存档）")
            time.sleep(delay)

        cache_path.write_text(json.dumps(cache, ensure_ascii=False), encoding="utf-8")
        filled = sum(1 for v in cache.values() if v)
        log(f"  完成：{len(cache)} 份合同，其中 {filled} 份有产品明细")
        return cache

    def contracts(self, date_from: str = "", date_to: str = "", page_size: int = 300,
                  log=print, max_pages: int = 100) -> list[dict]:
        """拉取合同列表，返回 [{列名: 值}]"""
        rows: list[dict] = []
        cols: list[str] = []
        page_index = 1
        total = None

        while page_index <= max_pages:
            params = {"pagesize": page_size, "pageindex": page_index}
            if date_from:
                params["dateQD_0"] = date_from
            if date_to:
                params["dateQD_1"] = date_to

            d = self.asp_call(PATH_CONTRACT, CMDKEY_CONTRACT, params)
            table = d.get("body", {}).get("source", {}).get("table")
            if not table:
                raise ErpError(f"合同接口返回结构不认识，顶层键：{list(d.get('body', {}).keys())}")

            cols = [c["id"] if isinstance(c, dict) else c for c in (table.get("cols") or [])]
            batch = table.get("rows") or []
            pg = table.get("page") or {}
            total = pg.get("recordcount", total)

            for r in batch:
                rows.append(dict(zip(cols, r)))
            if log:
                log(f"  第 {page_index} 页 +{len(batch)} 条，累计 {len(rows)}/{total}")

            if not batch or not pg.get("pagecount") or page_index >= pg["pagecount"]:
                break
            if total is not None and len(rows) >= total:
                break
            page_index += 1

        return rows


# --------------------------------------------------------------------------
# 转成报表能吃的标准表
# --------------------------------------------------------------------------

# ERP 字段 -> 报表列名
FIELD_MAP = [
    ("所属门店", "htfl"),      # 合同分类，实际存的就是门店名
    ("销售员",   "xsry"),      # 形如「万象城-韩亚辉」
    ("订单编号", "htid"),
    ("客户姓名", "khmc"),
    ("合同金额", "moneyall"),  # 优惠后金额
    ("下单日期", "dateQD"),    # 签订日期
    ("已回款",   "hkmoney"),
    ("合同主题", "title"),
]

HEADERS = [h for h, _ in FIELD_MAP]


def clean_person(raw: str) -> str:
    """「万象城-韩亚辉」-> 「韩亚辉」；没有前缀就原样返回"""
    s = core.cell_text(raw)
    for sep in ("-", "－", "—", "–"):
        if sep in s:
            head, _, tail = s.partition(sep)
            if tail.strip():
                return tail.strip()
    return s


def to_standard_table(rows: list[dict], strip_prefix: bool = True) -> list[list]:
    table = [list(HEADERS)]
    for r in rows:
        line = []
        for header, field in FIELD_MAP:
            v = r.get(field, "")
            if header == "销售员" and strip_prefix:
                v = clean_person(v)
            line.append(core.cell_text(v))
        table.append(line)
    return table


# --------------------------------------------------------------------------
# 保存
# --------------------------------------------------------------------------

def save_table(cfg: dict, table: list[list], tag: str = "") -> Path:
    import datetime as _dt

    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    out_dir = core.cfg_path(cfg, "原始数据目录")
    out_dir.mkdir(parents=True, exist_ok=True)
    stamp = _dt.datetime.now().strftime("%Y%m%d_%H%M")
    name = f"智邦ERP业绩_{tag + '_' if tag else ''}{stamp}.xlsx"
    target = out_dir / name

    # 同一区间重复导出会在目录里累积多份文件，报表读整个目录就会重复计数。
    # 保存前先把同区间的旧文件清掉，保证一个区间只有一份。
    removed = 0
    if tag:
        for old in out_dir.glob(f"智邦ERP业绩_{tag}_*.xlsx"):
            if old.resolve() == target.resolve():
                continue
            try:
                old.unlink()
                removed += 1
            except OSError:
                pass

    wb = Workbook()
    ws = wb.active
    ws.title = "智邦ERP业绩"
    for r in table:
        ws.append(r)

    head_fill = PatternFill("solid", fgColor="1F4E78")
    for c in ws[1]:
        c.font = Font(name="微软雅黑", size=10, bold=True, color="FFFFFF")
        c.fill = head_fill
        c.alignment = Alignment(horizontal="center")
    for i, h in enumerate(table[0], 1):
        width = max(len(str(h)) * 2.2, 10)
        for row in table[1:201]:
            if i - 1 < len(row):
                width = max(width, min(len(str(row[i - 1])) + 4, 42))
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.freeze_panes = "A2"
    wb.save(target)
    return target


# --------------------------------------------------------------------------

def load_erp_config(cfg: dict) -> dict:
    e = cfg.get("智邦ERP") or {}
    if not e.get("地址"):
        raise SystemExit(
            "config.json 里「智邦ERP.地址」是空的。\n"
            "填上你们 ERP 的网址，例如 http://tn166.kmdns.net:6088"
        )
    if not e.get("账号"):
        raise SystemExit("config.json 里「智邦ERP.账号」是空的。")
    if not e.get("密码"):
        raise SystemExit("config.json 里「智邦ERP.密码」是空的。")
    return e


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description="智邦国际 ERP 业绩抓取")
    ap.add_argument("--开始", help="签订日期起，如 2026-01-01")
    ap.add_argument("--结束", help="签订日期止，如 2026-12-31")
    ap.add_argument("--出报表", action="store_true", help="抓完立刻生成 Excel 报表")
    ap.add_argument("--探测", action="store_true", help="只测连通性和数据结构，不写文件")
    ap.add_argument("--保留前缀", action="store_true", help="销售员保留「万象城-」前缀")
    args = ap.parse_args(argv)

    cfg = core.load_config()
    e = load_erp_config(cfg)

    d_from = args.开始 or (e.get("日期范围") or {}).get("开始") or ""
    d_to = args.结束 or (e.get("日期范围") or {}).get("结束") or ""

    print(f"ERP 地址：{e['地址']}")
    print(f"账号：{e['账号']}")
    print(f"签订日期：{d_from or '不限'} ~ {d_to or '不限'}\n")

    erp = ZbErp(e["地址"], e["账号"], e["密码"], e.get("串号") or "zb-scraper-01")
    erp.login()

    if args.探测:
        print("\n--- 部门（门店）---")
        for d in erp.departments():
            path = d.get("名称全路径") or d.get("部门名称")
            if any(k in str(path) for k in ("贵和", "万象城", "中兴")):
                print(f"  id={d.get('部门标识')}  {path}")
        print("\n--- 合同接口 ---")
        rows = erp.contracts(d_from, d_to, page_size=5, max_pages=1, log=None)
        print(f"  取到 {len(rows)} 条样本，字段：{list(rows[0].keys()) if rows else '（无）'}")
        if rows:
            r = rows[0]
            print(f"  样例：门店={r.get('htfl')!r} 销售员={r.get('xsry')!r} "
                  f"金额={r.get('moneyall')!r} 签订日期={r.get('dateQD')!r}")
        return 0

    print("开始抓取合同...")
    rows = erp.contracts(d_from, d_to, page_size=int(e.get("每页条数") or 300))
    if not rows:
        raise SystemExit("没抓到任何合同。检查日期范围，或确认这个账号有销售数据的查看权限。")

    table = to_standard_table(rows, strip_prefix=not args.保留前缀)
    tag = f"{d_from}_{d_to}" if d_from or d_to else ""
    target = save_table(cfg, table, tag)
    print(f"\n已保存 {len(table) - 1} 条到：\n  {target}")

    # 顺手打印汇总，方便一眼核对
    stores: dict[str, dict[str, float]] = {}
    for line in table[1:]:
        store, person, amount = line[0], line[1], core.to_money(line[4])
        stores.setdefault(store, {}).setdefault(person, 0.0)
        stores[store][person] += amount
    for store, people in stores.items():
        total = sum(people.values())
        print(f"\n【{store}】{len(people)} 人 / {total:,.2f} 元")
        for name, amt in sorted(people.items(), key=lambda kv: -kv[1]):
            print(f"    {name:<10} {amt:>16,.2f} 元")

    if args.出报表:
        import subprocess
        print("\n" + "-" * 50 + "\n开始生成报表...\n")
        subprocess.run([sys.executable, str(core.BASE / "zbom_report.py")], check=False)
    else:
        print("\n下一步： python zbom_report.py")
    return 0


if __name__ == "__main__":
    try:
        sys.exit(main())
    except ErpError as exc:
        print(f"\n[抓取失败] {exc}", file=sys.stderr)
        sys.exit(1)

# -*- coding: utf-8 -*-
"""
志邦业绩汇总出报表
用法：python zbom_report.py
"""
from __future__ import annotations

import datetime as _dt
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

import zbom_core as core

# ---------------------------------------------------------------- 样式
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEAD_FILL = PatternFill("solid", fgColor="1F4E78")
HEAD_FONT = Font(name="微软雅黑", size=10, bold=True, color="FFFFFF")
TITLE_FONT = Font(name="微软雅黑", size=14, bold=True, color="1F4E78")
SUB_FONT = Font(name="微软雅黑", size=9, color="808080")
BODY_FONT = Font(name="微软雅黑", size=10)
BOLD_FONT = Font(name="微软雅黑", size=10, bold=True)
TOTAL_FILL = PatternFill("solid", fgColor="DDEBF7")

GOLD = PatternFill("solid", fgColor="FFF2CC")
SILVER = PatternFill("solid", fgColor="F2F2F2")
BRONZE = PatternFill("solid", fgColor="FCE4D6")
MONEY_FMT = '#,##0.00'
INT_FMT = '#,##0'
PCT_FMT = '0.0%'
CENTER = Alignment(horizontal="center", vertical="center")
LEFT = Alignment(horizontal="left", vertical="center")
RIGHT = Alignment(horizontal="right", vertical="center")


def _head(ws, row: int, headers: list[str], widths: list[int] | None = None):
    for i, h in enumerate(headers, 1):
        c = ws.cell(row=row, column=i, value=h)
        c.fill, c.font, c.alignment, c.border = HEAD_FILL, HEAD_FONT, CENTER, BORDER
    if widths:
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[row].height = 22
    ws.freeze_panes = ws.cell(row=row + 1, column=1)


def _title(ws, text: str, sub: str, span: int):
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max(span, 2))
    c = ws.cell(row=1, column=1, value=text)
    c.font, c.alignment = TITLE_FONT, Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[1].height = 26
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=max(span, 2))
    c2 = ws.cell(row=2, column=1, value=sub)
    c2.font, c2.alignment = SUB_FONT, Alignment(horizontal="left", vertical="center")


def rank_fill(rank: int):
    return {1: GOLD, 2: SILVER, 3: BRONZE}.get(rank)


def build(cfg: dict, records: list[dict], warnings: list[str], out_dir: Path | None = None) -> Path:
    stores = list(cfg.get("门店", {}).keys())
    summary = core.summarize(records, cfg)

    grand_total = sum(
        v["金额"] for s, ppl in summary.items() if s != "未匹配门店" for v in ppl.values()
    )
    grand_orders = sum(
        v["订单数"] for s, ppl in summary.items() if s != "未匹配门店" for v in ppl.values()
    )

    dr = cfg.get("数据源", {}).get("日期范围", {}) or {}
    ds = [r["日期"] for r in records if r.get("日期")]
    if dr.get("开始") or dr.get("结束"):
        period = f"{dr.get('开始') or '不限'} ~ {dr.get('结束') or '不限'}"
    elif ds:
        period = f"{min(ds)} ~ {max(ds)}"
    else:
        period = "不限"
    report_title = cfg.get("报表标题") or "门店业绩报表"
    stamp = _dt.datetime.now().strftime("%Y%m%d_%H%M")
    today = _dt.datetime.now().strftime("%Y-%m-%d %H:%M")

    # 任务按月折算：区间几个月，任务就乘几
    _ds = [r["日期"] for r in records if r.get("日期")]
    span = core.month_span(min(_ds).isoformat(), max(_ds).isoformat()) if _ds else 0.0
    tasks = {s: core.task_of(cfg, s, span) for s in stores} if span else {}

    wb = Workbook()

    # ---------------------------------------------------------- 总览
    ws = wb.active
    ws.title = "总览"
    _title(ws, f"{report_title} · 总览",
           f"数据区间：{period}　|　生成时间：{today}　|　门店 {len(stores)} 个"
           + (f"　|　任务按 {span} 个月折算" if span else ""), 12)
    _head(ws, 4, ["门店", "销售人数", "订单数", "业绩金额", "平均单值", "占总额比",
                  "最低任务", "达成", "标准任务", "达成", "门店销冠", "销冠金额"],
          [16, 10, 10, 16, 14, 11, 15, 10, 15, 10, 14, 16])

    def _pct_cell(ws_, row, col, val):
        c = ws_.cell(row=row, column=col, value=val)
        c.font, c.border = BODY_FONT, BORDER
        c.alignment = CENTER
        if val is None:
            c.value = "—"
            c.font = Font(name="微软雅黑", size=10, color="A6A6A6")
        else:
            c.number_format = PCT_FMT
            c.font = Font(name="微软雅黑", size=10, bold=True,
                          color="1B7F4B" if val >= 1 else ("C77700" if val >= 0.8 else "C00000"))
        return c

    r = 5
    store_rows = []
    for store in stores:
        ppl = summary.get(store, {})
        amt = sum(v["金额"] for v in ppl.values())
        cnt = sum(v["订单数"] for v in ppl.values())
        active = [n for n, v in ppl.items() if v["订单数"] > 0]
        top = next(iter(ppl.items()), None)
        store_rows.append((store, len(active), cnt, amt, top))
        t = tasks.get(store)
        vals = [
            store,
            len(active),
            cnt,
            amt,
            (amt / cnt) if cnt else 0,
            (amt / grand_total) if grand_total else 0,
            (t or {}).get("最低") or "—",
            None,                                    # 达成（下面单独写）
            (t or {}).get("标准") or "—",
            None,
            top[0] if top else "—",
            top[1]["金额"] if top else 0,
        ]
        for i, v in enumerate(vals, 1):
            if v is None:
                continue
            c = ws.cell(row=r, column=i, value=v)
            c.font, c.border = BODY_FONT, BORDER
            c.alignment = LEFT if i == 1 else (RIGHT if i in (4, 5, 7, 9, 12) else CENTER)
        for i in (4, 5, 12):
            ws.cell(row=r, column=i).number_format = MONEY_FMT
        for i in (7, 9):
            if isinstance(ws.cell(row=r, column=i).value, (int, float)):
                ws.cell(row=r, column=i).number_format = MONEY_FMT
        ws.cell(row=r, column=6).number_format = PCT_FMT
        lo = (t or {}).get("最低")
        std = (t or {}).get("标准")
        _pct_cell(ws, r, 8, (amt / lo) if lo else None)
        _pct_cell(ws, r, 10, (amt / std) if std else None)
        if top and top[0]:
            ws.cell(row=r, column=11).font = BOLD_FONT
        r += 1

    _lo_sum = sum((tasks.get(s) or {}).get("最低") or 0 for s in stores)
    _std_sum = sum((tasks.get(s) or {}).get("标准") or 0 for s in stores)
    for i, v in enumerate(["合计", len({n for s in stores for n in summary.get(s, {})}), grand_orders,
                           grand_total, (grand_total / grand_orders) if grand_orders else 0,
                           1 if grand_total else 0,
                           _lo_sum or "—", None, _std_sum or "—", None, "", ""], 1):
        if v is None:
            continue
        c = ws.cell(row=r, column=i, value=v)
        c.font, c.fill, c.border = BOLD_FONT, TOTAL_FILL, BORDER
        c.alignment = LEFT if i == 1 else (RIGHT if i in (4, 5, 7, 9, 12) else CENTER)
    for i in (4, 5, 12, 7, 9):
        if isinstance(ws.cell(row=r, column=i).value, (int, float)):
            ws.cell(row=r, column=i).number_format = MONEY_FMT
    ws.cell(row=r, column=6).number_format = PCT_FMT
    for col, ssum in ((8, _lo_sum), (10, _std_sum)):
        c = ws.cell(row=r, column=col, value=(grand_total / ssum) if ssum else "—")
        c.font, c.fill, c.border = BOLD_FONT, TOTAL_FILL, BORDER
        c.alignment = CENTER
        if ssum:
            c.number_format = PCT_FMT

    # 门店横向对比条形
    r += 3
    ws.cell(row=r, column=1, value="业绩对比").font = TITLE_FONT
    r += 1
    for store, n, cnt, amt, top in sorted(store_rows, key=lambda x: -x[3]):
        ws.cell(row=r, column=1, value=store).font = BODY_FONT
        ratio = (amt / grand_total) if grand_total else 0
        bar = "█" * max(1, int(ratio * 40)) if amt else ""
        c = ws.cell(row=r, column=2, value=bar)
        c.font = Font(name="微软雅黑", size=10, color="2E75B6")
        c2 = ws.cell(row=r, column=3, value=amt)
        c2.number_format = MONEY_FMT
        c2.font = BODY_FONT
        c3 = ws.cell(row=r, column=4, value=f"{ratio:.1%}")
        c3.font, c3.alignment = BODY_FONT, CENTER
        r += 1

    # 未匹配提示
    unmatched = summary.get("未匹配门店", {})
    if unmatched:
        r += 1
        c = ws.cell(row=r, column=1, value=f"⚠ 有 {sum(v['订单数'] for v in unmatched.values())} 条记录没能归到三个门店，已单独放在「未匹配门店」表，可在 config.json 的「门店」关键词里补充。")
        c.font = Font(name="微软雅黑", size=10, color="C00000")

    # ---------------------------------------------------------- 每个门店一张表
    for store in stores:
        _store_sheet(wb, store, summary.get(store, {}), period, today)

    # ---------------------------------------------------------- 全员排名
    ws4 = wb.create_sheet("全员排名")
    _title(ws4, f"全员业绩排名（{len(stores)} 店合并）", f"数据区间：{period}　|　按业绩金额降序", 6)
    _head(ws4, 4, ["排名", "销售员", "所属门店", "订单数", "业绩金额", "占总额比"],
          [8, 14, 14, 10, 16, 11])

    everyone = []
    for store in stores:
        for name, v in summary.get(store, {}).items():
            everyone.append((name, store, v))
    everyone.sort(key=lambda x: (-x[2]["金额"], x[1], x[0]))

    r = 5
    for i, (name, store, v) in enumerate(everyone, 1):
        vals = [i, name, store, v["订单数"], v["金额"], (v["金额"] / grand_total) if grand_total else 0]
        for j, val in enumerate(vals, 1):
            c = ws4.cell(row=r, column=j, value=val)
            c.font, c.border = BODY_FONT, BORDER
            c.alignment = CENTER if j in (1, 4, 6) else LEFT
        ws4.cell(row=r, column=5).number_format = MONEY_FMT
        ws4.cell(row=r, column=6).number_format = PCT_FMT
        f = rank_fill(i)
        if f:
            for j in range(1, 7):
                ws4.cell(row=r, column=j).fill = f
        r += 1

    for i, v in enumerate(["", "合计", f"{len(stores)} 个门店", grand_orders, grand_total, 1 if grand_total else 0], 1):
        c = ws4.cell(row=r, column=i, value=v)
        c.font, c.fill, c.border = BOLD_FONT, TOTAL_FILL, BORDER
        c.alignment = CENTER if i in (1, 4, 6) else LEFT
    ws4.cell(row=r, column=5).number_format = MONEY_FMT
    ws4.cell(row=r, column=6).number_format = PCT_FMT

    # ---------------------------------------------------------- 未匹配门店
    if unmatched:
        ws5 = wb.create_sheet("未匹配门店")
        _title(ws5, "未匹配到门店的记录", "请根据「门店原文」列，把对应关键词补进 config.json 的「门店」配置里", 6)
        _head(ws5, 4, ["销售员", "门店原文", "订单数", "金额", "来源文件", "客户举例"],
              [14, 24, 10, 14, 30, 20])
        r = 5
        for name, v in sorted(unmatched.items(), key=lambda kv: -kv[1]["金额"]):
            srcs = sorted({d["来源"] for d in v["明细"]})
            custs = [d["客户"] for d in v["明细"] if d["客户"]][:3]
            for j, val in enumerate([name, v["明细"][0]["门店原文"] if v["明细"] else "", v["订单数"],
                                     v["金额"], "、".join(srcs), "、".join(custs)], 1):
                c = ws5.cell(row=r, column=j, value=val)
                c.font, c.border = BODY_FONT, BORDER
                c.alignment = CENTER if j == 3 else LEFT
            ws5.cell(row=r, column=4).number_format = MONEY_FMT
            r += 1

    # ---------------------------------------------------------- 数据来源说明
    ws6 = wb.create_sheet("数据质量")
    _title(ws6, "数据质量与提示", "逐条检查这里，确认没有漏抓、没认错列", 3)
    _head(ws6, 4, ["级别", "说明"], [10, 110])
    r = 5
    notes = list(warnings)
    for name, stores in sorted(core.cross_store_people(records).items()):
        notes.append(f"「{name}」同时出现在 {'、'.join(stores)} —— 如果不是调店，多半是门店关键词认错了，请检查。")
    for s, ppl in summary.items():
        zero = [n for n, v in ppl.items() if v["订单数"] == 0]
        if s != "未匹配门店" and zero:
            notes.append(f"[{s}] 有 {len(zero)} 人本期无业绩记录：{'、'.join(zero[:20])}")
    if not records:
        notes.append("一条记录都没读到——请确认原始数据目录里有志邦导出的业绩表，且表里有「销售员」和「金额」两列。")
    if not notes:
        notes.append("没有发现问题。")
    for n in notes:
        lvl = "警告" if ("跳过" in n or "没找到" in n or "没有" in n) else "提示"
        c1 = ws6.cell(row=r, column=1, value=lvl)
        c1.font = Font(name="微软雅黑", size=10, bold=True,
                       color="C00000" if lvl == "警告" else "808080")
        c1.border, c1.alignment = BORDER, CENTER
        c2 = ws6.cell(row=r, column=2, value=n)
        c2.font, c2.border, c2.alignment = BODY_FONT, BORDER, LEFT
        ws6.row_dimensions[r].height = 18
        r += 1

    # ---------------------------------------------------------- 保存
    out_dir = out_dir or core.cfg_path(cfg, "目录")
    out_dir.mkdir(parents=True, exist_ok=True)
    prefix = cfg.get("输出", {}).get("文件名前缀", "志邦业绩报表")
    target = out_dir / f"{prefix}_{stamp}.xlsx"
    wb.save(target)
    return target


def _store_sheet(wb: Workbook, store: str, ppl: dict, period: str, today: str):
    ws = wb.create_sheet(store[:28])
    _title(ws, f"{store}　销售业绩明细", f"数据区间：{period}　|　生成时间：{today}", 9)

    total = sum(v["金额"] for v in ppl.values())
    orders = sum(v["订单数"] for v in ppl.values())

    _head(ws, 4, ["排名", "销售员", "订单数", "业绩金额", "占门店比", "平均单值", "最大单", "最早成交", "最晚成交"],
          [7, 14, 9, 15, 11, 13, 13, 13, 13])

    r = 5
    for name, v in ppl.items():
        vals = [v["排名"], name, v["订单数"], v["金额"], v["占比"], v["平均单值"], v["最大单"],
                v["最早"] or "", v["最晚"] or ""]
        for j, val in enumerate(vals, 1):
            c = ws.cell(row=r, column=j, value=val)
            c.font, c.border = BODY_FONT, BORDER
            c.alignment = LEFT if j == 2 else (RIGHT if j in (4, 6, 7) else CENTER)
        for j in (4, 6, 7):
            ws.cell(row=r, column=j).number_format = MONEY_FMT
        ws.cell(row=r, column=5).number_format = PCT_FMT
        f = rank_fill(v["排名"]) if v["订单数"] else None
        if f:
            for j in range(1, 10):
                ws.cell(row=r, column=j).fill = f
        if v["订单数"] == 0:
            ws.cell(row=r, column=2).font = Font(name="微软雅黑", size=10, color="A6A6A6")
        r += 1

    for j, val in enumerate(["", "门店合计", orders, total, 1 if total else 0,
                             (total / orders) if orders else 0, "", "", ""], 1):
        c = ws.cell(row=r, column=j, value=val)
        c.font, c.fill, c.border = BOLD_FONT, TOTAL_FILL, BORDER
        c.alignment = LEFT if j == 2 else (RIGHT if j in (4, 6, 7) else CENTER)
    for j in (4, 6):
        ws.cell(row=r, column=j).number_format = MONEY_FMT
    ws.cell(row=r, column=5).number_format = PCT_FMT

    # ---- 订单明细
    r += 2
    ws.cell(row=r, column=1, value=f"{store} 订单明细（共 {orders} 条）").font = TITLE_FONT
    r += 1
    for i, h in enumerate(["日期", "销售员", "客户", "订单号", "产品", "金额", "来源文件"], 1):
        c = ws.cell(row=r, column=i, value=h)
        c.fill, c.font, c.alignment, c.border = HEAD_FILL, HEAD_FONT, CENTER, BORDER
    r += 1

    detail = [d for v in ppl.values() for d in v["明细"]]
    detail.sort(key=lambda d: (d["日期"] or _dt.date.min, d["销售员"]), reverse=True)
    for d in detail:
        vals = [d["日期"] or "", d["销售员"], d["客户"], d["订单号"], d["产品"], d["金额"], d["来源"]]
        for j, val in enumerate(vals, 1):
            c = ws.cell(row=r, column=j, value=val)
            c.font, c.border = BODY_FONT, BORDER
            c.alignment = RIGHT if j == 6 else (CENTER if j in (1, 4) else LEFT)
        ws.cell(row=r, column=6).number_format = MONEY_FMT
        r += 1


def main(argv: list[str]) -> int:
    cfg = core.load_config()
    records, warnings = core.load_records(cfg)
    print(f"读取记录：{len(records)} 条")
    for w in warnings:
        print(f"  [警告] {w}")
    target = build(cfg, records, warnings)
    print(f"\n报表已生成：{target}")

    summary = core.summarize(records, cfg)
    for store in cfg.get("门店", {}):
        ppl = summary.get(store, {})
        amt = sum(v["金额"] for v in ppl.values())
        cnt = sum(v["订单数"] for v in ppl.values())
        print(f"\n【{store}】{cnt} 单 / {amt:,.2f} 元 / {len(ppl)} 人")
        for name, v in list(ppl.items())[:10]:
            print(f"   {v['排名']:>2}. {name:<8} {v['订单数']:>3} 单  {v['金额']:>14,.2f} 元  {v['占比']:>6.1%}")
        if len(ppl) > 10:
            print(f"   ... 其余 {len(ppl)-10} 人见报表")
    unmatched = summary.get("未匹配门店", {})
    if unmatched:
        print(f"\n【未匹配门店】{sum(v['订单数'] for v in unmatched.values())} 单，请补 config.json 关键词")
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
